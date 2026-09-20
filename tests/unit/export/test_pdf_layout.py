from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from bankocr.export.pdf_layout import PdfLayoutExcelExporter
from bankocr.parser.models import FieldValue, TransactionCandidate


def _make_two_page_pdf(path: Path) -> Path:
    document = pymupdf.open()
    first = document.new_page(width=600, height=800)
    first.insert_text((40, 60), "第一张原始页面", fontsize=18)
    second = document.new_page(width=800, height=600)
    second.insert_text((40, 60), "第二张原始页面", fontsize=18)
    document.save(path)
    document.close()
    return path


def _candidate(*, page_index: int, row_index: int) -> TransactionCandidate:
    return TransactionCandidate(
        page_index=page_index,
        row_index=row_index,
        parser_id="test:v1",
        fields={
            "amount": FieldValue("10.00", "10.00", confidence=0.99),
        },
    )


def test_pdf_layout_excel_contains_one_image_sheet_per_pdf_page_and_index(
    tmp_path: Path,
) -> None:
    source = _make_two_page_pdf(tmp_path / "statement.pdf")
    output = tmp_path / "statement.pdf-layout.xlsx"
    candidates = (
        _candidate(page_index=0, row_index=1),
        _candidate(page_index=1, row_index=2),
    )

    PdfLayoutExcelExporter().export(source, output, candidates)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["核对索引", "第001页", "第002页"]
    assert len(workbook["第001页"]._images) == 1
    assert len(workbook["第002页"]._images) == 1

    index = workbook["核对索引"]
    headers = [cell.value for cell in index[2]]
    assert "comparison_id" in headers
    assert "PDF页" in headers
    page_column = headers.index("PDF页") + 1
    assert index.cell(3, page_column).hyperlink.target == "#'第001页'!A1"
    assert index.cell(4, page_column).hyperlink.target == "#'第002页'!A1"

    document = pymupdf.open(source)
    ratios = [page.rect.width / page.rect.height for page in document]
    document.close()
    for page_number, ratio in enumerate(ratios, start=1):
        image = workbook[f"第{page_number:03d}页"]._images[0]
        assert image.width / image.height == pytest.approx(ratio, rel=0.02)


def test_pdf_layout_excel_preserves_blank_pages_and_rejects_source_overwrite(
    tmp_path: Path,
) -> None:
    source = tmp_path / "blank.pdf"
    document = pymupdf.open()
    document.new_page(width=300, height=300)
    document.new_page(width=300, height=300)
    document.save(source)
    document.close()

    output = tmp_path / "blank.pdf-layout.xlsx"
    PdfLayoutExcelExporter().export(source, output, ())
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["核对索引", "第001页", "第002页"]
    assert all(len(workbook[name]._images) == 1 for name in workbook.sheetnames[1:])

    with pytest.raises(ValueError, match="must not overwrite"):
        PdfLayoutExcelExporter().export(source, source, ())
