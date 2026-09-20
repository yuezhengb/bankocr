import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from bankocr.export.excel import ExcelExporter
from bankocr.export.manifest import ExportManifest
from bankocr.parser.models import FieldValue, TransactionCandidate


def test_export_manifest_records_file_hashes_and_sizes(tmp_path: Path) -> None:
    excel = tmp_path / "statement.xlsx"
    review = tmp_path / "statement.review.xlsx"
    searchable = tmp_path / "statement.searchable.pdf"
    comparison = tmp_path / "statement.comparison.pdf"
    excel.write_bytes(b"excel")
    review.write_bytes(b"review")
    searchable.write_bytes(b"pdf")
    comparison.write_bytes(b"comparison")

    manifest = ExportManifest.create(
        run_id=7,
        source_sha256="a" * 64,
        unresolved_count=2,
        outputs={
            "excel": excel,
            "review_excel": review,
            "searchable_pdf": searchable,
            "comparison_pdf": comparison,
        },
    )
    destination = tmp_path / "statement.export-manifest.json"
    manifest.write(destination)

    payload = json.loads(destination.read_text(encoding="utf-8"))
    assert payload["run_id"] == 7
    assert payload["source_sha256"] == "a" * 64
    assert payload["unresolved_count"] == 2
    assert payload["outputs"]["excel"]["size"] == 5
    assert payload["outputs"]["excel"]["sha256"] == hashlib.sha256(b"excel").hexdigest()


def test_transactions_sheet_puts_business_columns_before_technical_columns(tmp_path: Path) -> None:
    candidate = TransactionCandidate(
        page_index=0,
        row_index=0,
        parser_id="test:v1",
        fields={
            "transaction_date": FieldValue("2026-08-01", "2026-08-01"),
            "transaction_amount": FieldValue("10.00", "10.00"),
            "balance": FieldValue("10.00", "10.00"),
        },
    )
    output = tmp_path / "transactions.xlsx"

    ExcelExporter().export((candidate,), output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    headers = tuple(next(workbook["Transactions"].iter_rows(values_only=True)))
    assert headers[:9] == (
        "date",
        "income",
        "expense",
        "balance",
        "summary",
        "counterparty_account",
        "counterparty_name",
        "page",
        "status",
    )
    assert headers.index("transaction_id") > headers.index("status")
