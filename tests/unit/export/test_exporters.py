from pathlib import Path
from datetime import date

import pymupdf
import pytest
from openpyxl import load_workbook

from bankocr.domain.coordinates import Point
from bankocr.domain.text_block import SourceType, TextBlock
from bankocr.export.excel import ExcelExporter
from bankocr.export.searchable_pdf import SearchablePdfExporter
from bankocr.parser.models import FieldValue, TransactionCandidate
from bankocr.validation.engine import (
    IssueSeverity,
    ValidationIssue,
    ValidationReport,
)
from bankocr.validation.risk import ValidationRuleState


def _candidate() -> TransactionCandidate:
    return TransactionCandidate(
        page_index=0,
        row_index=3,
        parser_id="template:test:v1",
        fields={
            "date": FieldValue("2026/08/01", "2026-08-01", confidence=0.99),
            "amount": FieldValue("1O.00", "10.00", secondary_text="10.00", confidence=0.72),
        },
    )


def _block(text: str) -> TextBlock:
    return TextBlock(
        text=text,
        raw_text=text,
        polygon=(Point(20, 20), Point(80, 20), Point(80, 35), Point(20, 35)),
        confidence=0.9,
        source_type=SourceType.OCR,
        page_index=0,
        engine_id="test",
    )


def test_excel_export_keeps_raw_suggested_and_final_columns_separate(tmp_path: Path) -> None:
    output = tmp_path / "transactions.xlsx"

    ExcelExporter().export((_candidate(),), output)

    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True, data_only=True)
    sheet = workbook["Transactions"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    values = tuple(next(sheet.iter_rows(min_row=2, values_only=True)))
    by_header = dict(zip(headers, values))

    assert by_header["amount.raw"] == "1O.00"
    assert by_header["amount.secondary"] == 10
    assert by_header["amount.suggested"] == 10
    assert by_header["amount.final"] is None


def test_excel_export_includes_processing_report_and_source_index(tmp_path: Path) -> None:
    output = tmp_path / "transactions.xlsx"
    report = ValidationReport(
        checked_rows=1,
        issues=(
            ValidationIssue(
                code="invalid_amount",
                message="invalid",
                severity=IssueSeverity.ERROR,
                row_index=3,
                field_name="amount",
            ),
        ),
        final_balances=(None,),
    )

    ExcelExporter().export(
        (_candidate(),),
        output,
        source_file="statement.pdf",
        validation_reports={0: report},
    )

    from openpyxl import load_workbook

    workbook = load_workbook(output, read_only=True, data_only=True)
    assert set(workbook.sheetnames) == {"Transactions", "Processing Report", "Source Index"}
    report_rows = list(workbook["Processing Report"].iter_rows(values_only=True))
    report_map = {row[0]: row[1] for row in report_rows if row and row[0]}
    assert report_map["source_file"] == "statement.pdf"
    assert report_map["validation_issue_count"] == 1

    source_headers = tuple(next(workbook["Source Index"].iter_rows(values_only=True)))
    source_values = tuple(next(workbook["Source Index"].iter_rows(min_row=2, values_only=True)))
    source_map = dict(zip(source_headers, source_values))
    assert source_map["field"] == "amount"
    assert source_map["raw_text"] == "1O.00"
    assert source_map["source_block_indices"] is None


def test_excel_export_includes_cross_check_id_and_source_location(tmp_path: Path) -> None:
    output = tmp_path / "cross-check.xlsx"
    located = TransactionCandidate(
        page_index=0,
        row_index=1,
        parser_id="test:v1",
        fields={
            "amount": FieldValue(
                "10.00",
                "10.00",
                source_block_indices=(0,),
            )
        },
    )
    unlocated = TransactionCandidate(
        page_index=0,
        row_index=2,
        parser_id="test:v1",
        fields={"amount": FieldValue("20.00", "20.00")},
    )

    ExcelExporter().export(
        (unlocated, located),
        output,
        blocks_by_page={0: (_block("10.00"),)},
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    sheet = workbook["Transactions"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    by_row = {row["row_index"]: row for row in rows}

    assert by_row[1]["comparison_id"] == "T0001"
    assert by_row[1]["source_page"] == 1
    assert by_row[1]["source_bbox"] == "20,20,80,35"
    assert by_row[1]["source_location_status"] == "located"
    assert by_row[2]["comparison_id"] == "T0002"
    assert by_row[2]["source_bbox"] is None
    assert by_row[2]["source_location_status"] == "unlocated"


def test_excel_export_keeps_marked_duplicate_visible_but_without_transaction_id(tmp_path: Path) -> None:
    output = tmp_path / "duplicate.xlsx"
    original = _candidate()
    duplicate = TransactionCandidate(
        page_index=0,
        row_index=4,
        parser_id=original.parser_id,
        fields=original.fields,
    ).mark_duplicate((0, 3))

    ExcelExporter().export((original, duplicate), output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    sheet = workbook["Transactions"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    duplicate_row = dict(zip(headers, rows[1]))
    assert duplicate_row["status"] == "duplicate"
    assert duplicate_row["duplicate_of"] == "0:3"
    assert duplicate_row["transaction_id"] is None


def test_excel_source_index_uses_matching_field_rule_state(tmp_path: Path) -> None:
    output = tmp_path / "rule-state.xlsx"
    report = ValidationReport(
        checked_rows=1,
        issues=(
            ValidationIssue(
                code="balance_mismatch",
                message="different row",
                severity=IssueSeverity.CRITICAL,
                row_index=99,
            ),
            ValidationIssue(
                code="invalid_amount",
                message="amount needs review",
                severity=IssueSeverity.ERROR,
                row_index=3,
                field_name="amount",
                rule_state=ValidationRuleState.INDETERMINATE,
            ),
        ),
        final_balances=(None,),
    )

    ExcelExporter().export((_candidate(),), output, validation_reports={0: report})

    workbook = load_workbook(output, read_only=True, data_only=True)
    headers = tuple(next(workbook["Source Index"].iter_rows(values_only=True)))
    values = tuple(next(workbook["Source Index"].iter_rows(min_row=2, values_only=True)))
    row = dict(zip(headers, values))
    assert row["rule_state"] == "indeterminate"


def test_excel_export_writes_normalized_amounts_and_dates_as_excel_types(tmp_path: Path) -> None:
    output = tmp_path / "typed.xlsx"
    candidate = TransactionCandidate(
        page_index=0,
        row_index=0,
        parser_id="test:v1",
        fields={
            "transaction_date": FieldValue("2026/08/01", "2026-08-01"),
            "transaction_amount": FieldValue("1,000.00", "1000.00"),
        },
    )

    ExcelExporter().export((candidate,), output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    headers = tuple(next(workbook["Transactions"].iter_rows(values_only=True)))
    values = tuple(next(workbook["Transactions"].iter_rows(min_row=2, values_only=True)))
    by_header = dict(zip(headers, values))
    assert isinstance(by_header["transaction_date.suggested"], date)
    assert by_header["transaction_date.suggested"].date() == date(2026, 8, 1)
    assert isinstance(by_header["transaction_amount.suggested"], (int, float))
    assert by_header["transaction_amount.suggested"] == 1000


def test_searchable_pdf_export_does_not_overwrite_input_and_adds_extractable_text(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    output = tmp_path / "searchable.pdf"
    document = pymupdf.open()
    page = document.new_page(width=200, height=100)
    page.insert_text((20, 50), "original", fontsize=10)
    document.save(source)
    document.close()

    SearchablePdfExporter().export(source, output, {0: (_block("10.00"),)})

    assert source.exists()
    assert output.exists()
    searchable = pymupdf.open(output)
    try:
        assert "10.00" in searchable[0].get_text()
    finally:
        searchable.close()


def test_searchable_pdf_export_keeps_cjk_text_extractable_with_explicit_font(
    tmp_path: Path,
) -> None:
    font = Path(r"C:\Windows\Fonts\simhei.ttf")
    if not font.is_file():
        pytest.skip("Windows CJK font is not available on this host")
    source = tmp_path / "source-cjk.pdf"
    output = tmp_path / "searchable-cjk.pdf"
    document = pymupdf.open()
    document.new_page(width=200, height=100)
    document.save(source)
    document.close()

    block = TextBlock(
        text="交易摘要",
        raw_text="交易摘要",
        polygon=(Point(20, 20), Point(100, 20), Point(100, 35), Point(20, 35)),
        confidence=0.9,
        source_type=SourceType.OCR,
        page_index=0,
        engine_id="test",
    )
    SearchablePdfExporter(font_file=font).export(source, output, {0: (block,)})

    searchable = pymupdf.open(output)
    try:
        assert "交易摘要" in searchable[0].get_text()
    finally:
        searchable.close()
