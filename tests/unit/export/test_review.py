from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from bankocr.export.review import ReviewExcelExporter
from bankocr.parser.models import FieldValue, TransactionCandidate
from bankocr.validation.engine import IssueSeverity, ValidationIssue, ValidationReport


def test_review_export_contains_issue_reason_and_provenance(tmp_path: Path) -> None:
    candidate = TransactionCandidate(
        page_index=0,
        row_index=2,
        parser_id="template:test:v1",
        fields={
            "balance": FieldValue(
                raw_text="102.00",
                suggested_text="102.00",
                secondary_text="100.00",
                confidence=0.64,
                source_block_indices=(4, 5),
            )
        },
    )
    report = ValidationReport(
        checked_rows=1,
        issues=(
            ValidationIssue(
                code="balance_mismatch",
                message="expected 100.00, got 102.00",
                severity=IssueSeverity.CRITICAL,
                row_index=2,
                field_name="balance",
            ),
        ),
        final_balances=(Decimal("102.00"),),
    )
    output = tmp_path / "review.xlsx"

    ReviewExcelExporter().export(
        (candidate,),
        {0: report},
        output,
        source_file="statement.pdf",
    )

    workbook = load_workbook(output, read_only=True, data_only=True)
    sheet = workbook["Review Queue"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    values = tuple(next(sheet.iter_rows(min_row=2, values_only=True)))
    row = dict(zip(headers, values))
    assert row["source_file"] == "statement.pdf"
    assert row["code"] == "balance_mismatch"
    assert row["severity"] == "critical"
    assert row["raw_text"] == "102.00"
    assert row["secondary_text"] == "100.00"
    assert row["source_block_indices"] == "4,5"
