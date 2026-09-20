"""Traceable Excel export with raw, suggested, and final columns."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Mapping, Sequence

from bankocr.parser.models import CandidateStatus, TransactionCandidate
from bankocr.parser.normalization import normalize_amount, normalize_date
from bankocr.transaction.builder import TransactionBuilder
from bankocr.validation.engine import ValidationReport
from bankocr.validation.risk import ValidationRuleState

from .source_geometry import candidate_bbox, candidate_sort_key, comparison_ids, format_bbox


class ExcelExporter:
    def export(
        self,
        candidates: tuple[TransactionCandidate, ...] | list[TransactionCandidate],
        output: str | Path,
        *,
        source_file: str | None = None,
        validation_reports: Mapping[int, ValidationReport] | None = None,
        metadata: Mapping[str, object] | None = None,
        blocks_by_page: Mapping[int, Sequence[object]] | None = None,
    ) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        candidate_values = tuple(sorted(candidates, key=candidate_sort_key))
        comparison_id_map = comparison_ids(candidate_values)
        transaction_records = {
            (record.candidate.page_index, record.candidate.row_index): record
            for record in TransactionBuilder().build(candidate_values)
        }
        output_path = Path(output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        field_names = sorted({name for candidate in candidate_values for name, _ in candidate.fields})
        headers = [
            "date",
            "income",
            "expense",
            "balance",
            "summary",
            "counterparty_account",
            "counterparty_name",
            "page",
            "status",
            "comparison_id",
            "source_page",
            "source_bbox",
            "source_location_status",
            "duplicate_of",
            "validation_status",
            "transaction_id",
            "statement_stream_id",
            "row",
            "page_index",
            "row_index",
            "parser_id",
        ]
        for name in field_names:
            headers.extend(
                (f"{name}.raw", f"{name}.secondary", f"{name}.suggested", f"{name}.final")
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Transactions"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for candidate in candidate_values:
            report = (validation_reports or {}).get(candidate.page_index)
            record = transaction_records.get((candidate.page_index, candidate.row_index))
            row = [
                _excel_value(alias, _effective_field_text(candidate, aliases))
                for alias, aliases in _CANONICAL_FIELD_ALIASES.items()
            ]
            row.extend(
                (
                    candidate.page_index + 1,
                    candidate.status.value,
                    comparison_id_map[(candidate.page_index, candidate.row_index)],
                    candidate.page_index + 1,
                    format_bbox(candidate_bbox(candidate, blocks_by_page)),
                    "located" if candidate_bbox(candidate, blocks_by_page) is not None else "unlocated",
                    _duplicate_reference(candidate),
                    report.status.value if report is not None and report.status is not None else "unknown",
                    record.transaction_id if record is not None else None,
                    record.statement_stream_id if record is not None else None,
                    candidate.row_index,
                    candidate.page_index,
                    candidate.row_index,
                    candidate.parser_id,
                )
            )
            fields = dict(candidate.fields)
            for name in field_names:
                value = fields.get(name)
                row.extend(
                    (
                        value.raw_text if value else None,
                        _excel_value(name, value.secondary_text) if value else None,
                        _excel_value(name, value.suggested_text) if value else None,
                        _excel_value(name, value.final_text) if value else None,
                    )
                )
            sheet.append(row)
            if report is not None and report.status is not None and report.status.value in {"review", "page_review"}:
                for cell in sheet[sheet.max_row]:
                    cell.fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(40, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width

        reports = validation_reports or {}
        report_sheet = workbook.create_sheet("Processing Report")
        report_sheet.append(("metric", "value"))
        report_rows = (
            ("source_file", source_file or ""),
            ("page_count", len({candidate.page_index for candidate in candidate_values} | set(reports))),
            ("transaction_count", len(candidate_values)),
            (
                "accepted_count",
                sum(candidate.status.value == "accepted" for candidate in candidate_values),
            ),
            (
                "rejected_count",
                sum(candidate.status.value == "rejected" for candidate in candidate_values),
            ),
            (
                "duplicate_count",
                sum(candidate.status is CandidateStatus.DUPLICATE for candidate in candidate_values),
            ),
            (
                "needs_review_count",
                sum(candidate.status.value == "unreviewed" for candidate in candidate_values),
            ),
            ("validation_issue_count", sum(len(report.issues) for report in reports.values())),
            (
                "confirmed_failure_count",
                sum(
                    sum(issue.rule_state.value == "fail" for issue in report.issues)
                    for report in reports.values()
                ),
            ),
            (
                "indeterminate_count",
                sum(
                    sum(issue.rule_state.value == "indeterminate" for issue in report.issues)
                    for report in reports.values()
                ),
            ),
            (
                "critical_issue_count",
                sum(sum(issue.severity.value == "critical" for issue in report.issues) for report in reports.values()),
            ),
            (
                "validation_statuses",
                ", ".join(
                    f"{page}:{report.status.value if report.status is not None else 'unknown'}"
                    for page, report in sorted(reports.items())
                ),
            ),
        )
        for row in (*report_rows, *((str(key), value) for key, value in (metadata or {}).items())):
            report_sheet.append(row)
        report_sheet.freeze_panes = "A2"
        report_sheet.column_dimensions["A"].width = 28
        report_sheet.column_dimensions["B"].width = 48

        source_sheet = workbook.create_sheet("Source Index")
        source_sheet.append(
            (
                "transaction_id",
                "comparison_id",
                "page_index",
                "row_index",
                "field",
                "raw_text",
                "secondary_text",
                "suggested_text",
                "final_text",
                "confidence",
                "source_block_indices",
                "source_spans",
                "source_bbox",
                "validation_status",
                "rule_state",
                "status",
            )
        )
        for candidate in candidate_values:
            record = transaction_records.get((candidate.page_index, candidate.row_index))
            for name, value in candidate.fields:
                source_sheet.append(
                    (
                        record.transaction_id if record is not None else None,
                        comparison_id_map[(candidate.page_index, candidate.row_index)],
                        candidate.page_index,
                        candidate.row_index,
                        name,
                        value.raw_text,
                        value.secondary_text,
                        value.suggested_text,
                        value.final_text,
                        value.confidence,
                        ",".join(str(index) for index in value.source_block_indices),
                        _source_spans_json(value.source_spans),
                        _source_bbox(candidate.page_index, value.source_block_indices, blocks_by_page),
                        (validation_reports or {}).get(candidate.page_index).status.value
                        if (validation_reports or {}).get(candidate.page_index) is not None
                        and (validation_reports or {}).get(candidate.page_index).status is not None
                        else "unknown",
                        _rule_state_for(
                            (validation_reports or {}).get(candidate.page_index),
                            candidate.row_index,
                            name,
                        ),
                        candidate.status.value,
                    )
                )
        source_sheet.freeze_panes = "A2"
        source_sheet.auto_filter.ref = source_sheet.dimensions
        for column_cells in source_sheet.columns:
            width = min(40, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            source_sheet.column_dimensions[column_cells[0].column_letter].width = width
        workbook.save(output_path)
        return output_path


_AMOUNT_FIELDS = frozenset(
    {
        "amount",
        "balance",
        "online_balance",
        "opening_balance",
        "closing_balance",
        "income",
        "expense",
    }
)

_CANONICAL_FIELD_ALIASES: Mapping[str, tuple[str, ...]] = {
    "date": ("transaction_date", "accounting_date", "date"),
    "income": ("income_amount", "income"),
    "expense": ("expense_amount", "expense"),
    "balance": ("balance", "online_balance", "closing_balance"),
    "summary": ("summary", "transaction_summary", "description", "customer_summary"),
    "counterparty_account": ("counterparty_account", "account_number", "account"),
    "counterparty_name": ("counterparty_name", "counterparty_info", "name"),
}


def _effective_field_text(candidate: TransactionCandidate, names: tuple[str, ...]) -> str | None:
    fields = dict(candidate.fields)
    for name in names:
        value = fields.get(name)
        if value is None:
            continue
        return value.final_text or value.suggested_text or value.raw_text
    return None


def _duplicate_reference(candidate: TransactionCandidate) -> str | None:
    if candidate.duplicate_of is None:
        return None
    page_index, row_index = candidate.duplicate_of
    return f"{page_index}:{row_index}"


def _source_bbox(
    page_index: int,
    source_indices: tuple[int, ...],
    blocks_by_page: Mapping[int, Sequence[object]] | None,
) -> str | None:
    if not blocks_by_page or not source_indices:
        return None
    blocks = blocks_by_page.get(page_index, ())
    boxes = [getattr(blocks[index], "bbox") for index in source_indices if 0 <= index < len(blocks)]
    if not boxes:
        return None
    x0 = min(box[0] for box in boxes)
    y0 = min(box[1] for box in boxes)
    x1 = max(box[2] for box in boxes)
    y1 = max(box[3] for box in boxes)
    return f"{x0:g},{y0:g},{x1:g},{y1:g}"


def _source_spans_json(source_spans: object) -> str | None:
    spans = tuple(source_spans or ())
    if not spans:
        return None
    return json.dumps(
        [
            {
                "page_index": span.page_index,
                "text_block_id": span.text_block_id,
                "polygon": [(point.x, point.y) for point in span.polygon],
            }
            for span in spans
        ],
        ensure_ascii=False,
        separators=(",", ":"),
    )


def _rule_state_for(
    report: ValidationReport | None,
    row_index: int,
    field_name: str,
) -> str:
    if report is None:
        return "pass"
    exact = [
        issue
        for issue in report.issues
        if issue.row_index == row_index and issue.field_name == field_name
    ]
    row_level = [
        issue
        for issue in report.issues
        if issue.row_index == row_index and issue.field_name is None
    ]
    matching = exact or row_level
    if not matching:
        return "pass"
    priority = {
        ValidationRuleState.FAIL: 0,
        ValidationRuleState.INDETERMINATE: 1,
        ValidationRuleState.PASS: 2,
        ValidationRuleState.NOT_APPLICABLE: 3,
    }
    return min(matching, key=lambda issue: priority[issue.rule_state]).rule_state.value


def _excel_value(field_name: str, text: str | None) -> object:
    if text is None:
        return None
    normalized_name = field_name.casefold()
    if normalized_name == "date" or normalized_name.endswith("_date"):
        parsed_date = normalize_date(text)
        return parsed_date if parsed_date is not None else text
    if normalized_name in _AMOUNT_FIELDS or normalized_name.endswith("_amount"):
        parsed_amount = normalize_amount(text)
        return parsed_amount if parsed_amount is not None else text
    return text
