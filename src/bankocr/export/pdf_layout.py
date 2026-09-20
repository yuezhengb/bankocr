"""Excel workbook that preserves each source PDF page as an image."""

from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf

from bankocr.parser.models import TransactionCandidate

from .source_geometry import candidate_sort_key, comparison_ids


class PdfLayoutExcelExporter:
    """Export the original PDF artwork into a page-per-sheet Excel workbook.

    This exporter deliberately does not place OCR text or comparison markers on
    the page images.  The structured Excel and comparison PDF remain separate
    artifacts for those purposes.
    """

    def __init__(self, *, dpi: int = 150) -> None:
        if dpi <= 0:
            raise ValueError("PDF layout Excel dpi must be positive")
        self.dpi = dpi

    def export(
        self,
        source: str | Path,
        output: str | Path,
        candidates: Sequence[TransactionCandidate],
    ) -> Path:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.utils import get_column_letter

        source_path = Path(source).expanduser().resolve()
        output_path = Path(output).expanduser().resolve()
        if source_path == output_path:
            raise ValueError("PDF layout Excel output must not overwrite the source")
        output_path.parent.mkdir(parents=True, exist_ok=True)

        ordered_candidates = tuple(sorted(candidates, key=candidate_sort_key))
        comparison_id_map = comparison_ids(ordered_candidates)
        workbook = Workbook()
        index = workbook.active
        index.title = "核对索引"
        _configure_index(index, source_path.name)
        _append_index_rows(index, ordered_candidates, comparison_id_map)

        document = pymupdf.open(str(source_path))
        try:
            with TemporaryDirectory(prefix="bankocr-pdf-layout-", dir=output_path.parent) as temp_dir:
                for page_index, page in enumerate(document):
                    pixmap = page.get_pixmap(dpi=self.dpi, alpha=False)
                    image_path = Path(temp_dir) / f"page-{page_index + 1:04d}.png"
                    pixmap.save(str(image_path))

                    sheet = workbook.create_sheet(f"第{page_index + 1:03d}页")
                    sheet.sheet_view.showGridLines = False
                    sheet.page_setup.orientation = (
                        "landscape" if page.rect.width > page.rect.height else "portrait"
                    )
                    sheet.sheet_properties.pageSetUpPr.fitToPage = True
                    sheet.page_setup.fitToWidth = 1
                    sheet.page_setup.fitToHeight = 1
                    sheet.page_margins.left = 0
                    sheet.page_margins.right = 0
                    sheet.page_margins.top = 0
                    sheet.page_margins.bottom = 0
                    sheet.page_margins.header = 0
                    sheet.page_margins.footer = 0

                    image = ExcelImage(str(image_path))
                    image.width = pixmap.width
                    image.height = pixmap.height
                    sheet.add_image(image, "A1")
                    _configure_print_grid(sheet, pixmap.width, pixmap.height, get_column_letter)

                temporary_output = Path(temp_dir) / output_path.name
                workbook.save(temporary_output)
                temporary_output.replace(output_path)
        finally:
            document.close()
        return output_path


def _configure_index(sheet, source_name: str) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    sheet["A1"] = (
        f"{source_name}：本页用于和原 PDF 视觉对照；搜索、统计和整理请使用普通 Excel。"
    )
    sheet["A1"].font = Font(bold=True, color="1F4E78")
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    headers = (
        "comparison_id",
        "PDF页",
        "流水行号",
        "状态",
        "解析器",
        "字段摘要",
        "页面工作表",
    )
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=2, column=column, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    widths = (18, 12, 14, 14, 24, 70, 18)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"


def _append_index_rows(
    sheet,
    candidates: Sequence[TransactionCandidate],
    comparison_id_map: Mapping[tuple[int, int], str],
) -> None:
    for row, candidate in enumerate(candidates, start=3):
        comparison_id = comparison_id_map[(candidate.page_index, candidate.row_index)]
        page_sheet = f"第{candidate.page_index + 1:03d}页"
        page_cell = sheet.cell(row=row, column=2, value=candidate.page_index + 1)
        page_cell.hyperlink = f"#'{page_sheet}'!A1"
        page_cell.style = "Hyperlink"
        sheet.cell(row=row, column=1, value=comparison_id)
        sheet.cell(row=row, column=3, value=candidate.row_index)
        sheet.cell(row=row, column=4, value=candidate.status.value)
        sheet.cell(row=row, column=5, value=candidate.parser_id)
        sheet.cell(row=row, column=6, value=_field_summary(candidate))
        page_link = sheet.cell(row=row, column=7, value=page_sheet)
        page_link.hyperlink = f"#'{page_sheet}'!A1"
        page_link.style = "Hyperlink"
    last_row = max(2, sheet.max_row)
    sheet.auto_filter.ref = f"A2:G{last_row}"


def _field_summary(candidate: TransactionCandidate) -> str:
    parts = []
    for name, value in candidate.fields:
        text = value.final_text or value.suggested_text or value.secondary_text or value.raw_text
        if text:
            parts.append(f"{name}={text}")
    return "；".join(parts)[:32_000]


def _configure_print_grid(
    sheet,
    width: int,
    height: int,
    get_column_letter: Callable[[int], str],
) -> None:
    # These dimensions are only a print/viewing grid for the floating image;
    # they do not claim that the page has been converted into editable cells.
    column_count = max(1, (width + 69) // 70)
    row_count = max(1, (height + 19) // 20)
    for column in range(1, column_count + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    for row in range(1, row_count + 1):
        sheet.row_dimensions[row].height = 15
    sheet.print_area = f"A1:{get_column_letter(column_count)}{row_count}"
